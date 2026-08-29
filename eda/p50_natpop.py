#!/usr/bin/env python
"""Phase 50 — Korea's age-band population vector, and the anchors that it is right.

WHY THIS EXISTS. Every survey matrix in this project is reduced to a contact
matrix by `symmetrise(C, pop)`: T = N_a C[a,a'], symmetrised. `pop` has always
come from `results_p26.json["population"]`, and `p26_matrix.py:46` says what
that is -- "N_a is Seoul's registered population". For the Seoul arm that is
correct. For the NATIONAL arm it is not: 1,987 respondents drawn from the whole
country are being reweighted onto Seoul's age structure, and the result is
neither a Seoul matrix nor a national one. The advisor's 2026-08-24 letter is
blunt about the consequence -- 不修的話那一臂稱不上全國臂.

This phase builds the vector that arm should have been using. It does NOT touch
`regpop.parquet`, `foreign.parquet`, `results_p20.json` or `results_p26.json`;
it writes `results_p50.json` and nothing else.

THE DATA WAS ALREADY ON DISK. No download was needed and none was made. The
행정안전부 registration CSV that `dl_regpop.py` fetches has always been the
NATIONAL 읍면동 table (`xlsStats=3`, ~3,800 rows, 17 시도); `p20_regpop.py:201`
throws away everything outside Seoul with `if not code.startswith("11")`. The
법무부 foreigner spreadsheet is likewise national, and `p20b_foreign.py:166`
drops the other sixteen 시도. The monthly 시군구 anchor `p20b` calibrates onto,
source A (`moj_gu_month.csv`), is national too -- `p20b_foreign.py:187` filters
it to Seoul because p20b builds a Seoul product, not because the file is one.
All three filters are correct for what those phases build. This phase reads the
same three files without them.

THE ANCHORS, and why the foreigner half now has one at both months.

  TIER 1, 주민등록, BIT-EXACT. Restricting the national build to codes beginning
  `11` must reproduce, band by band and exactly, `results_p26.json["population"]
  [ym]` minus the Seoul foreigner margin. This is a strong anchor and not a
  tautology: it goes through a different reader, skips p20's crosswalk, ABSORBED
  map and 제-ordinal normalisation entirely, and still has to land on the same
  sixteen integers. It caught nothing, which is the point -- it says p20's
  dong-level surgery does not move the age margin.

  A second route on the same half: the dong-level single-year columns summed
  against the published 시도 roll-up totals, which are different cells of the
  same file. Both months are +0.

  TIER 2, 등록외국인, BIT-EXACT AT BOTH MONTHS. 202312 is a quarter end and is
  read directly. 202402 is not -- and 2024-03, the quarter that would have
  covered it, is published as a legacy BIFF .xls that openpyxl cannot read
  (`dl_foreign.py:387`, `p20b_foreign.py:44`), so the dong-level composition has
  to be interpolated across 2023-12..2024-06. What does NOT have to be
  interpolated is the level: source A publishes 시군구 totals every month for
  the whole country, 202402 included. So this phase does what p20b does for
  Seoul -- interpolate the composition, then rescale each 시군구 block onto A --
  and the Seoul restriction reproduces `foreign.parquet`'s Seoul margin to
  2.2e-16 instead of the 3.8e-02 that plain interpolation left.

  The rescale is written as a ratio to A's OWN linear interpolation, not to A's
  level, so at a quarter end the factor is exactly 1 and 202312 cannot move.
  A only ever supplies within-quarter movement.

  This replaces a declaration that fired. The previous run declared "interpolated
  months reported and capped at 1e-3 relative per band", observed 3.79e-02, and
  marked 202402 `established=False`. The tolerance was NOT widened; the anchor
  the data already supported was built instead. `superseded_declaration` in the
  results file keeps the old rule and the number it fired at.

TWO DEFECTS THIS ROUND FOUND, both in the foreigner half.

  1. 세종특별자치시 WAS BEING DROPPED ENTIRELY. It is the one 시도 with no 시군구
     layer: its dong rows carry an empty 시군구 cell, and `''` is in p20b's
     ROLLUP set, so the Seoul-shaped filter `if gu in ROLLUP: continue` deleted
     a whole province. 5,786 people at 202312 -- exactly the gap against source
     A, which is how it was found. Harmless in p20b, which never sees 세종.
  2. 2024-01-01 부천시 re-established 소사/오정/원미 구, so the gu key at 2023-12
     is not the gu key at 2024-06. The three are merged back to 부천시 in every
     source (`GU_MERGE`) rather than interpolated from zero.

WHAT THE TWO SOURCES AGREE ON, which is worth stating as a finding. After those
two repairs, the quarterly dong spreadsheet and the monthly 시군구 file -- two
separate 법무부 publications, different portals, different shapes -- agree cell
for cell at the quarter ends: 250 of 250 시군구 at 202312, both summing to
1,348,626, and 249 of 250 at 202406, the exception being 세종 by 18 people. That
is why the ratio rescale is a no-op at a quarter end by measurement and not
merely by construction, and it is the whole of the residual at 202402: the
rescaled national total is 1,354,960 against A's 1,354,954, and the +6 is 세종's
18 carried in at w = 1/3. Nothing else in the country disagrees at all.

WHY THE FOREIGNER COMPONENT IS WORTH THE TROUBLE AT ALL. It is 2.6% of Seoul's
denominator and it is not age-flat -- 46,249 of Seoul's 252,765 registered
foreigners in 202312 are 20-24, a band where they are 8.3% of the total. A
national vector that silently omitted them would not be comparable to the Seoul
vector it is replacing.

THE FIELD THAT MAKES A SILENTLY-WRONG VECTOR VISIBLE. `share_ratio` is the
16-band Seoul share divided by the Korea share. If a downstream phase ever
reverts to the Seoul vector by accident, this field is how it is noticed: the
ratio runs from about 0.83 at 10-14 to about 1.26 at 25-29, so the two vectors
are not interchangeable and the file says so in a form a reader can check.

    python eda/p50_natpop.py
"""
import csv
import io
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import numpy as np
import openpyxl
import pandas as pd
from common import AGES, AGE_LABEL
from p20_regpop import AGE_BIN, N_AGE_COLS, age_bin
from p20b_foreign import BAND_TO_BIN, DONG_COL, ROLLUP, midx, norm_age
from paths import DERIVED, RAW, ROOT

YMS = (202312, 202402)              # the two survey months
REGDIR = f"{RAW}/raw/regpop"
MOJDIR = f"{RAW}/raw/foreign"
CODE_RE = re.compile(r"\((\d{10})\)")
SEOUL_TOL = 1e-9                    # tier 2: relative, per band, BOTH months

# The 시도 label is not stable across sources or across 2023-2024: the dong
# spreadsheet says 강원특별자치도 and 전북특별자치도, source A still says 강원도
# and 전라북도, and A's 경기 arrives as 경기, 경기도 and 경기도 with a trailing
# ideographic space. Mapped explicitly and aborted on, never defaulted -- an
# unrecognised 시도 that silently fell through would drop a province, which is
# defect 1 above happening a second time.
SIDO_CANON = {
    "서울특별시": "서울", "부산광역시": "부산", "대구광역시": "대구",
    "인천광역시": "인천", "광주광역시": "광주", "대전광역시": "대전",
    "울산광역시": "울산", "세종특별자치시": "세종",
    "경기": "경기", "경기도": "경기",
    "강원도": "강원", "강원특별자치도": "강원",
    "충청북도": "충북", "충청남도": "충남",
    "전라북도": "전북", "전북특별자치도": "전북", "전라남도": "전남",
    "경상북도": "경북", "경상남도": "경남", "제주특별자치도": "제주",
}
NO_GU_SIDO = {"세종"}               # no 시군구 layer; it keys on itself
GU_MERGE = {                        # 2024-01-01: 부천시 re-established its 구
    ("경기", "부천시 소사구"): ("경기", "부천시"),
    ("경기", "부천시 오정구"): ("경기", "부천시"),
    ("경기", "부천시 원미구"): ("경기", "부천시"),
}

out = {}


def say(s=""):
    print(s, flush=True)


def canon_sido(s):
    s = s.replace("　", " ").strip()
    if s not in SIDO_CANON:
        raise SystemExit(f"unknown 시도 {s!r} -- extend SIDO_CANON; a 시도 that "
                         f"falls through this map is a province dropped in "
                         f"silence")
    return SIDO_CANON[s]


def gu_key(sido_raw, gu_raw):
    """One (시도, 시군구) key, canonical across all three files."""
    sd = canon_sido(sido_raw)
    gu = gu_raw.replace("　", " ").strip()
    if sd in NO_GU_SIDO:
        return (sd, sd)
    if not gu:
        raise SystemExit(f"{sd} has an empty 시군구 but is not in NO_GU_SIDO -- "
                         f"either a new 시도 lost its gu layer or the row is "
                         f"malformed")
    return GU_MERGE.get((sd, gu), (sd, gu))


# ------------------------------------------------------------------ 주민등록
def read_regpop_national(ym):
    """The national 읍면동 registration table -> (national, seoul) band vectors.

    p20_regpop.read_month with the Seoul filter removed and the dong-name
    machinery omitted, because an age margin does not depend on dong names.
    Roll-up rows (시도 totals, 시군구 totals) are skipped so nothing is counted
    twice; the 계 block is skipped because it is redundant with 남 + 여, which
    is what dl_regpop verifies on download.
    """
    path = f"{REGDIR}/regpop_{ym}.csv"
    text = open(path, "rb").read().decode("cp949")
    rows = list(csv.reader(io.StringIO(text)))
    header, body = rows[0], [r for r in rows[1:] if r and r[0].strip()]
    off = {}
    for i, ko in enumerate(("계", "남", "여")):
        base = 1 + i * (2 + N_AGE_COLS)
        y, m = divmod(ym, 100)
        want = f"{y}년{m:02d}월_{ko}_총인구수"
        if header[base].strip() != want:
            raise SystemExit(f"regpop_{ym}.csv: column {base} is "
                             f"{header[base]!r}, expected {want!r}")
        off[ko] = base + 2

    def num(c):
        c = c.strip().replace(",", "")
        return int(c) if c and c != "-" else 0

    nat = {a: 0 for a in AGES}
    seo = {a: 0 for a in AGES}
    n_sido = n_dong = n_seoul_dong = 0
    sido_total = 0
    for r in body:
        m = CODE_RE.search(r[0])
        if not m:
            raise SystemExit(f"regpop_{ym}.csv: no code in label {r[0]!r}")
        code = m.group(1)
        if code.endswith("00000000"):          # 시도 roll-up
            n_sido += 1
            sido_total += num(r[1])
            continue
        if code.endswith("00000"):             # 시군구 roll-up
            continue
        n_dong += 1
        is_seoul = code.startswith("11")
        n_seoul_dong += is_seoul
        for ko in ("남", "여"):
            base = off[ko]
            for a in range(N_AGE_COLS):
                v = num(r[base + a])
                if v:
                    b = AGE_BIN[a]
                    nat[b] += v
                    if is_seoul:
                        seo[b] += v
    return (np.array([nat[a] for a in AGES], float),
            np.array([seo[a] for a in AGES], float),
            dict(n_sido=n_sido, n_dong=n_dong, n_seoul_dong=n_seoul_dong,
                 sido_total=sido_total))


def regpop_second_route(ym):
    """An independent recount of the national total, csv + dict, no numpy.

    The invariant says a headline number needs a second implementation. This
    walks the 시도 roll-up rows and sums their published 총인구수 -- a
    completely different set of cells from the dong-level single-year columns
    the first route adds up. If the two agree, no dong was dropped and no roll-up
    was double-counted.
    """
    path = f"{REGDIR}/regpop_{ym}.csv"
    text = open(path, "rb").read().decode("cp949")
    tot = 0
    for r in csv.reader(io.StringIO(text)):
        if not r or not r[0].strip():
            continue
        m = CODE_RE.search(r[0])
        if m and m.group(1).endswith("00000000"):
            c = r[1].strip().replace(",", "")
            tot += int(c) if c and c != "-" else 0
    return tot


# --------------------------------------------------------------------- 법무부
def read_moj_national(path):
    """One 법무부 quarterly spreadsheet -> {(시도, 시군구): 16-band vector}.

    The header handling is p20b_foreign.read_moj's, verbatim in intent, because
    the layout drifts across quarters in ways that were each found by a file
    that failed verification rather than by reading metadata. What is removed is
    the `시도 startswith 서울` filter; what is repaired is the gu filter, which
    in p20b deletes every row with an empty 시군구 and so would delete all of
    세종 (defect 1 in the module docstring).
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    head, body = None, []
    for row in ws.iter_rows(values_only=True):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if head is None:
            if "성별" in cells and "총합계" in cells and "시도" in cells:
                head = cells
            continue
        if cells and cells[0]:
            body.append(list(row))
    wb.close()
    if head is None:
        raise SystemExit(f"{os.path.basename(path)}: no header row")

    idx = {n: head.index(n) for n in ("시군구", "성별", "총합계") if n in head}
    for n in DONG_COL:
        if n in head:
            idx["행정동"] = head.index(n)
            break
    sido = [i for i, h in enumerate(head) if h == "시도"]
    idx["시도"] = max([i for i in sido if i < idx["시군구"]] or sido)
    band = {i: BAND_TO_BIN[norm_age(head[i])]
            for i in range(idx["총합계"] + 1, len(head))
            if head[i] and norm_age(head[i]) in BAND_TO_BIN}
    if sorted(set(band.values())) != sorted(AGE_LABEL):
        raise SystemExit(f"{os.path.basename(path)}: bands do not cover 16 bins")

    def num(r, i):
        if i >= len(r) or r[i] is None or str(r[i]).strip() in ("", "-"):
            return 0
        return int(float(str(r[i]).replace(",", "")))

    def txt(r, i):
        return str(r[i]).strip() if i < len(r) and r[i] is not None else ""

    has_dong = "행정동" in idx
    cells = {}
    for r in body:
        sd = txt(r, idx["시도"])
        if sd in ROLLUP:
            continue
        gu = txt(r, idx["시군구"])
        if gu and gu in ROLLUP:            # a roll-up LABEL, not an empty cell
            continue
        if has_dong:
            dong = txt(r, idx["행정동"])
            if dong in ROLLUP or not dong:
                continue
        if txt(r, idx["성별"]) not in ("남성", "여성"):
            continue                       # the 총계 roll-up
        k = gu_key(sd, gu)
        v = cells.setdefault(k, np.zeros(len(AGES)))
        for i, a in band.items():
            n = num(r, i)
            if n:
                v[AGES.index(a)] += n
    return cells


def read_month_gu_national(ym):
    """Source A: month x 시군구 registered-foreigner totals, ALL of Korea.

    p20b_foreign.read_month_gu with the `startswith("서울")` filter removed.
    This is the file that makes 202402 anchorable: it is monthly, so the level
    of a mid-quarter month does not have to be interpolated even when the
    dong-level composition does.
    """
    path = f"{MOJDIR}/moj_gu_month.csv"
    text = open(path, "rb").read().decode("cp949")
    rows = list(csv.reader(io.StringIO(text)))
    head = [c.strip() for c in rows[0]]
    c = {n: head.index(n) for n in ("년", "월", "시도", "시군구", "등록외국인수")}
    y, m = divmod(ym, 100)
    tot = {}
    for r in rows[1:]:
        if not r or not r[0].strip():
            continue
        if int(r[c["년"]]) != y or int(r[c["월"]]) != m:
            continue
        k = gu_key(r[c["시도"]], r[c["시군구"]])
        tot[k] = tot.get(k, 0) + int(r[c["등록외국인수"]].replace(",", ""))
    if not tot:
        raise SystemExit(f"source A has no rows for {ym}")
    return tot


def bands(cells, seoul_only=False):
    v = np.zeros(len(AGES))
    for (sd, _), x in cells.items():
        if seoul_only and sd != "서울":
            continue
        v += x
    return v


def moj_for(ym):
    """Foreigner band vectors at `ym`: (national, seoul, meta).

    A quarter end is read directly and cross-checked against source A. Any other
    month has its dong-level composition interpolated linearly on the month
    index between the surrounding quarters -- p20b_foreign.py:33's first step --
    and is then rescaled, 시군구 by 시군구, by A's ratio to its own linear
    interpolation. At a quarter end that ratio is 1 by construction, so 202312
    cannot move; between quarters it is exactly the within-quarter movement the
    interpolation cannot see.
    """
    have = sorted(int(f.split("_")[-1].split(".")[0])
                  for f in os.listdir(MOJDIR)
                  if f.startswith("moj_dong_age_"))
    A_ym = read_month_gu_national(ym)

    if ym in have:
        cells = read_moj_national(f"{MOJDIR}/moj_dong_age_{ym}.xlsx")
        gaps = {k: float(cells[k].sum() - A_ym[k])
                for k in set(cells) & set(A_ym) if cells[k].sum() != A_ym[k]}
        seoul_gap = max([abs(v) for k, v in gaps.items() if k[0] == "서울"] or [0.])
        if seoul_gap:
            raise SystemExit(f"{ym}: the quarterly dong file and source A "
                             f"disagree about Seoul by {seoul_gap:,.0f} -- p20b "
                             f"rescales onto A, so they must agree here")
        meta = dict(mode="exact", quarter=ym,
                    sidos=len({k[0] for k in cells}),
                    gus=len(cells),
                    sejong=float(cells[("세종", "세종")].sum()),
                    cross_source=dict(
                        matched=len(set(cells) & set(A_ym)),
                        identical=len(set(cells) & set(A_ym)) - len(gaps),
                        max_abs_gap=max([abs(v) for v in gaps.values()] or [0.]),
                        only_in_dong=sorted("|".join(k) for k in
                                            set(cells) - set(A_ym)),
                        only_in_A=sorted("|".join(k) for k in
                                         set(A_ym) - set(cells)),
                        dong_total=float(bands(cells).sum()),
                        A_total=float(sum(A_ym.values()))))
        return bands(cells), bands(cells, True), meta

    lo = max(q for q in have if q < ym)
    hi = min(q for q in have if q > ym)
    c0, c1 = (read_moj_national(f"{MOJDIR}/moj_dong_age_{q}.xlsx")
              for q in (lo, hi))
    A_lo, A_hi = read_month_gu_national(lo), read_month_gu_national(hi)
    w = (midx(ym) - midx(lo)) / (midx(hi) - midx(lo))

    keys = sorted(set(c0) | set(c1))
    z = np.zeros(len(AGES))
    interp = {k: c0.get(k, z) + w * (c1.get(k, z) - c0.get(k, z)) for k in keys}
    before = float(bands(interp).sum())

    # The two 법무부 publications at the bracketing quarter ends. This is what
    # makes the ratio rescale a no-op at a quarter end by MEASUREMENT and not
    # merely by construction, and it is where the residual against A's own
    # national total comes from: whatever the two sources disagree about at a
    # quarter is carried into the interpolation, weighted by w.
    agree = {}
    for q, cq, Aq in ((lo, c0, A_lo), (hi, c1, A_hi)):
        gaps = {k: float(cq[k].sum() - Aq[k])
                for k in set(cq) & set(Aq) if cq[k].sum() != Aq[k]}
        agree[str(q)] = dict(
            matched=len(set(cq) & set(Aq)),
            identical=len(set(cq) & set(Aq)) - len(gaps),
            max_abs_gap=max([abs(v) for v in gaps.values()] or [0.]),
            differing=sorted(f"{'|'.join(k)}:{v:+.0f}" for k, v in gaps.items()),
            only_in_dong=sorted("|".join(k) for k in set(cq) - set(Aq)),
            only_in_A=sorted("|".join(k) for k in set(Aq) - set(cq)))

    resc, unanchored = {}, []
    for k in keys:
        base = A_lo.get(k, 0) + w * (A_hi.get(k, 0) - A_lo.get(k, 0))
        if k in A_ym and base > 0:
            resc[k] = interp[k] * (A_ym[k] / base)
        else:
            resc[k] = interp[k]
            unanchored.append(k)
    if any(k[0] == "서울" for k in unanchored):
        raise SystemExit(f"{ym}: a Seoul 시군구 has no source-A anchor "
                         f"{[k for k in unanchored if k[0] == '서울']} -- the "
                         f"tier-2 anchor rests on every Seoul gu being anchored")
    meta = dict(mode="interpolated+A", lo=lo, hi=hi, weight=float(w),
                sidos=len({k[0] for k in keys}), gus=len(keys),
                sejong=float(resc[("세종", "세종")].sum()),
                anchor=dict(
                    anchored=len(keys) - len(unanchored),
                    unanchored=sorted("|".join(k) for k in unanchored),
                    unanchored_people=float(sum(interp[k].sum()
                                                for k in unanchored)),
                    national_before=before,
                    national_after=float(bands(resc).sum()),
                    A_total=float(sum(A_ym.values())),
                    only_in_A=sorted("|".join(k) for k in set(A_ym) - set(keys)),
                    quarter_agreement=agree))
    return bands(resc), bands(resc, True), meta


def main():
    # ======================================================================= run
    say("=== 50.1 the national registration table, and the Seoul anchor ===")
    p26 = json.load(open(f"{ROOT}/eda/results_p26.json"))
    fo_par = pd.read_parquet(f"{DERIVED}/foreign.parquet")

    vectors, checks = {}, []
    for ym in YMS:
        nat_r, seo_r, meta = read_regpop_national(ym)
        total2 = regpop_second_route(ym)
        say(f"  {ym}: {meta['n_sido']} 시도, {meta['n_dong']:,} dong "
            f"({meta['n_seoul_dong']} of them Seoul)")
        say(f"        national total {nat_r.sum():>14,.0f}  "
            f"(시도 roll-up route {total2:,}, delta {nat_r.sum() - total2:+,.0f})")
        assert nat_r.sum() == total2, (
            f"{ym}: the dong-level sum and the 시도 roll-up sum disagree by "
            f"{nat_r.sum() - total2:+,.0f} -- a dong was dropped or double-counted")

        # --- TIER 1: the Seoul restriction, bit-exact -------------------------
        pub = np.array(p26["population"][str(ym)], float)
        fo_seoul = (fo_par[fo_par.ym == ym].groupby("age")["pop"].sum()
                    .reindex(AGES, fill_value=0).to_numpy(float))
        want = pub - fo_seoul
        dmax = float(np.abs(seo_r - want).max())
        say(f"        tier 1  Seoul restriction vs p26 population minus "
            f"foreigners: max abs diff {dmax:.1f}")
        assert dmax == 0.0, (
            f"{ym}: the national build restricted to Seoul is not p26's registered "
            f"component (max {dmax:,.0f}); p20's Seoul path moves the age margin")

        # --- TIER 2: the foreigner component, now anchored at both months -----
        nat_f, seo_f, fmeta = moj_for(ym)
        nz = fo_seoul > 0
        rel = np.abs(seo_f - fo_seoul)[nz] / fo_seoul[nz]
        if fmeta["mode"] == "exact":
            cs = fmeta["cross_source"]
            say(f"        tier 2  foreigners exact ({fmeta['gus']} 시군구, "
                f"{fmeta['sidos']} 시도); vs source A: {cs['identical']} of "
                f"{cs['matched']} 시군구 identical, dong {cs['dong_total']:,.0f} "
                f"vs A {cs['A_total']:,.0f}")
        else:
            an = fmeta["anchor"]
            say(f"        tier 2  foreigners {fmeta['mode']} "
                f"({fmeta['lo']}->{fmeta['hi']}, w={fmeta['weight']:.3f}); "
                f"{an['anchored']} of {fmeta['gus']} 시군구 rescaled onto A, "
                f"national {an['national_before']:,.0f} -> "
                f"{an['national_after']:,.0f} (A says {an['A_total']:,.0f})")
            for q, ag in an["quarter_agreement"].items():
                say(f"                quarter {q}: dong file vs source A, "
                    f"{ag['identical']} of {ag['matched']} 시군구 identical"
                    + (f", differing {ag['differing']}" if ag['differing'] else ""))
        say(f"                Seoul restriction vs foreign.parquet: max per-band "
            f"rel dev {rel.max():.2e} (tol {SEOUL_TOL:.0e})")
        assert rel.max() < SEOUL_TOL, (
            f"{ym}: the Seoul restriction of the national foreigner build deviates "
            f"from foreign.parquet's Seoul margin by {rel.max():.2e}")
        established = True

        nat = nat_r + nat_f
        seoul = pub
        share_nat = nat / nat.sum()
        share_seo = seoul / seoul.sum()
        vectors[str(ym)] = dict(
            national=nat.tolist(), seoul=seoul.tolist(),
            national_regpop=nat_r.tolist(), national_foreign=nat_f.tolist(),
            national_total=float(nat.sum()), seoul_total=float(seoul.sum()),
            share_national=share_nat.tolist(), share_seoul=share_seo.tolist(),
            share_ratio=(share_seo / share_nat).tolist(),
            foreigner_mode=fmeta, established=established,
            seoul_restriction_max_rel=float(rel.max()),
            max_band_share_gap_pp=float(100 * np.abs(share_seo - share_nat).max()))
        checks.append(dict(ym=ym, tier1_max_abs=dmax, established=established,
                           tier2_max_rel=float(rel.max()),
                           tier2_mode=fmeta["mode"],
                           regpop_two_routes_agree=True,
                           n_sido=meta["n_sido"], n_dong=meta["n_dong"],
                           n_foreign_gu=fmeta["gus"]))

    say("\n=== 50.2 the two vectors are not interchangeable ===")
    say(f"  {'band':>7} {'Korea':>13} {'Seoul':>12} {'Korea %':>8} {'Seoul %':>8} "
        f"{'ratio':>7}")
    v = vectors[str(YMS[0])]
    for i, a in enumerate(AGES):
        say(f"  {AGE_LABEL[a]:>7} {v['national'][i]:>13,.0f} "
            f"{v['seoul'][i]:>12,.0f} {100 * v['share_national'][i]:>7.2f}% "
            f"{100 * v['share_seoul'][i]:>7.2f}% {v['share_ratio'][i]:>7.3f}")
    rr = np.array(v["share_ratio"])
    say(f"  max |ratio - 1| = {np.abs(rr - 1).max():.3f} at "
        f"{AGE_LABEL[AGES[int(np.argmax(np.abs(rr - 1)))]]}")
    say(f"  largest band-share gap = {v['max_band_share_gap_pp']:.2f} pp")

    out["declaration"] = dict(
        months=list(YMS),
        sources=dict(regpop="행정안전부 주민등록 읍면동 (national, dl_regpop.py)",
                     foreign="법무부 등록외국인 시도x시군구x행정동 (national)",
                     foreign_monthly="법무부/data.go.kr 15100022 월 x 시군구 "
                                     "등록외국인수 (national, source A)"),
        tier1="Seoul restriction of the national regpop build == "
              "results_p26.json population minus foreign.parquet's Seoul margin, "
              "bit-exact, abort otherwise",
        tier2=f"foreigner component: quarter ends read directly, mid-quarter months "
              f"interpolated on the dong composition and rescaled per 시군구 by "
              f"source A's ratio to its own interpolation; the Seoul restriction "
              f"must reproduce foreign.parquet's Seoul margin to "
              f"{SEOUL_TOL:.0e} relative per band at BOTH months, abort otherwise",
        writes_only="results_p50.json",
        touches_nothing=["regpop.parquet", "foreign.parquet", "results_p20.json",
                         "results_p26.json"],
    )
    out["superseded_declaration"] = dict(
        rule="foreigner component: exact quarters bit-exact; interpolated months "
             "reported and capped at 0.001 relative per band",
        observed_at_202402=0.03791184283042286,
        what_it_measured="plain linear interpolation against p20b's interpolation "
                         "plus its Seoul monthly gu rescale",
        response="the tolerance was not widened; source A turned out to be national, "
                 "so the same rescale was built for the whole country and the "
                 "deviation is now 0 to float precision",
        also_found="세종특별자치시 was being dropped by the inherited gu filter "
                   "(5,786 people at 202312) and 부천시's 2024-01-01 구 split was "
                   "not being merged",
    )
    out["checks"] = checks
    out["vectors"] = vectors
    out["max_share_ratio_deviation"] = float(np.abs(rr - 1).max())

    with open(f"{ROOT}/eda/results_p50.json", "w") as fh:
        json.dump(out, fh, indent=1, ensure_ascii=False)
    say(f"\nwrote {ROOT}/eda/results_p50.json")


if __name__ == "__main__":
    main()
